"""
framework/reporting/report_manager.py
=======================================
Generates execution summary reports from ExecutionResult.

Output formats:
- JSON summary (machine-readable)
- Excel report (human-readable, with row-level details)
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import List

from config.environment import get_env
from config.constants import FRAMEWORK_NAME, FRAMEWORK_VERSION
from models.execution_result import ExecutionResult, RowResult
from framework.core.logger import get_logger

_log = get_logger()


class ReportManager:
    """
    Generates execution summary reports.

    Usage:
        rm = ReportManager()
        rm.generate(result)
    """

    def __init__(self) -> None:
        self._output_dir: Path = get_env().output_dir

    # ------------------------------------------------------------------ #
    # Public API                                                           #
    # ------------------------------------------------------------------ #

    def generate(self, result: ExecutionResult) -> str:
        """
        Generate all report formats for a completed execution.

        Args:
            result: Finalized ExecutionResult.

        Returns:
            Path to the JSON summary report.
        """
        result.finalize()
        json_path = self._write_json(result)
        self._write_excel(result)
        self._print_summary(result)
        return json_path

    # ------------------------------------------------------------------ #
    # JSON Report                                                          #
    # ------------------------------------------------------------------ #

    def _write_json(self, result: ExecutionResult) -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{result.workflow_name}_{timestamp}.json"
        path = self._output_dir / filename

        report_data = {
            "framework": FRAMEWORK_NAME,
            "version": FRAMEWORK_VERSION,
            "generated_at": datetime.now().isoformat(),
            "summary": result.summary(),
            "rows": [
                {
                    "row_id": r.row_id,
                    "status": r.status,
                    "duration_seconds": round(r.duration_seconds, 3),
                    "processed_at": r.processed_at.isoformat() if r.processed_at else None,
                    "error": r.error_message,
                    "screenshot": r.screenshot_path,
                }
                for r in result.row_results
            ],
        }

        try:
            with path.open("w", encoding="utf-8") as fh:
                json.dump(report_data, fh, indent=2, ensure_ascii=False)
            _log.info("JSON report saved: {path}", path=path.name)
            result.report_file_path = str(path)
        except Exception as exc:
            _log.error("Failed to write JSON report: {exc}", exc=exc)

        return str(path)

    # ------------------------------------------------------------------ #
    # Excel Report                                                         #
    # ------------------------------------------------------------------ #

    def _write_excel(self, result: ExecutionResult) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            _log.warning("openpyxl not installed. Skipping Excel report generation.")
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{result.workflow_name}_{timestamp}.xlsx"
        path = self._output_dir / filename

        wb = openpyxl.Workbook()

        # ---- Summary Sheet -----------------------------------------
        ws_summary = wb.active
        ws_summary.title = "Summary"

        summary = result.summary()
        ws_summary.append(["Field", "Value"])
        for k, v in summary.items():
            ws_summary.append([k, str(v)])

        # ---- Row Details Sheet -------------------------------------
        ws_rows = wb.create_sheet("Row Details")
        ws_rows.append([
            "Row ID", "Status", "Duration (s)", "Processed At", "Error", "Screenshot"
        ])

        for row in result.row_results:
            ws_rows.append([
                row.row_id,
                row.status,
                round(row.duration_seconds, 3),
                row.processed_at.isoformat() if row.processed_at else "",
                row.error_message or "",
                row.screenshot_path or "",
            ])

        try:
            wb.save(str(path))
            _log.info("Excel report saved: {path}", path=path.name)
        except Exception as exc:
            _log.error("Failed to write Excel report: {exc}", exc=exc)

    # ------------------------------------------------------------------ #
    # Console Summary                                                      #
    # ------------------------------------------------------------------ #

    def _print_summary(self, result: ExecutionResult) -> None:
        summary = result.summary()
        _log.info("=" * 60)
        _log.info("EXECUTION COMPLETE: {workflow}", workflow=summary["workflow"])
        _log.info("Execution ID : {id}", id=summary["execution_id"])
        _log.info("Duration     : {dur}s", dur=summary["duration_seconds"])
        _log.info("Total Rows   : {total}", total=summary["total"])
        _log.info("✅ Success   : {s}", s=summary["success"])
        _log.info("❌ Failure   : {f}", f=summary["failure"])
        _log.info("⏭  Skipped   : {sk}", sk=summary["skipped"])
        _log.info("Success Rate : {rate}%", rate=summary["success_rate_pct"])
        _log.info("=" * 60)
